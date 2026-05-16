from __future__ import annotations

import html
import mimetypes
import os
import threading
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal
from uuid import uuid4

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, PlainTextResponse
from docx import Document as DocxDocument
from pydantic import BaseModel
from langchain_core.messages import HumanMessage
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .config import get_config
from .graph import workflow

BASE_DIR = Path(__file__).resolve().parents[2]
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "test_output"

DEFAULT_TASK = "Analyze the customs declaration PDF and return structured output."

app = FastAPI(title="Docflow API")

app.add_middleware(
  CORSMiddleware,
  allow_origins=["*"],
  allow_credentials=True,
  allow_methods=["*"],
  allow_headers=["*"],
)


class UploadResponse(BaseModel):
  fileId: str


class StartJobRequest(BaseModel):
  fileId: str


class StartJobResponse(BaseModel):
  jobId: str


class JobStatusResponse(BaseModel):
  jobId: str
  status: Literal["queued", "running", "review_required", "completed", "failed"]
  step: str | None = None
  error: str | None = None


class OcrSummaryResponse(BaseModel):
  text: str
  reviewLevel: Literal["none", "warning", "error"] | None = None


class ContinueGenerationResponse(BaseModel):
  jobId: str
  status: Literal["running"]


class Artifact(BaseModel):
  id: str
  name: str
  type: str
  mimeType: str
  previewUrl: str | None = None
  downloadUrl: str


class DocumentItem(BaseModel):
  id: str
  name: str
  type: str
  mimeType: str
  date: str
  source: str
  jobId: str
  previewUrl: str | None = None
  downloadUrl: str


@dataclass
class JobRecord:
  file_id: str
  file_path: Path
  output_dir: Path
  status: Literal["queued", "running", "review_required", "completed", "failed"]
  step: str
  error: str | None = None
  result: dict | None = None
  documents: list[Path] = field(default_factory=list)
  created_at: float = field(default_factory=time.time)
  completed_at: float | None = None


FILES: dict[str, Path] = {}
JOBS: dict[str, JobRecord] = {}

HIDDEN_DOCUMENT_NAMES = {
  "ocr_output.txt",
  "rpa_mapping.json",
  "rpa_raw_response.txt",
}


def _safe_filename(name: str) -> str:
  return Path(name).name


def _original_upload_name(file_id: str, file_path: Path) -> str:
  prefixed_name = file_path.name
  prefix = f"{file_id}_"
  if prefixed_name.startswith(prefix):
    return prefixed_name[len(prefix):]
  return prefixed_name


def _resolve_path(path_str: str) -> Path:
  path = Path(path_str)
  if path.is_absolute():
    return path
  return BASE_DIR / path


def _detect_type(path: Path) -> tuple[str, str]:
  ext = path.suffix.lower()
  if ext in {".xlsx", ".xls"}:
    return "excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
  if ext in {".docx", ".doc"}:
    return "word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
  if ext == ".pdf":
    return "pdf", "application/pdf"
  if ext in {".png", ".jpg", ".jpeg"}:
    return "image", "image/jpeg"
  if ext in {".txt", ".md"}:
    return "text", "text/plain"

  mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
  return "other", mime


def _is_visible_document(path: Path) -> bool:
  return path.name not in HIDDEN_DOCUMENT_NAMES


def _document_timestamp(path: Path) -> float:
  stat = path.stat()
  created_at = getattr(stat, "st_birthtime", None)
  if created_at is not None and created_at > 0:
    return created_at
  return stat.st_mtime


def _preview_url(job_id: str, path: Path, doc_type: str) -> str | None:
  download_url = f"/documents/{job_id}/{path.name}"
  if doc_type in {"pdf", "text", "image"}:
    return download_url
  if doc_type in {"excel", "word"}:
    return f"{download_url}/preview"
  return None


def _find_document_path(job_id: str, filename: str) -> Path:
  safe_name = _safe_filename(filename)
  job = JOBS.get(job_id)

  if job:
    for item in _documents_for_job(job):
      if item.name == safe_name:
        return item

  candidate = OUTPUT_DIR / job_id / safe_name
  if candidate.exists() and _is_visible_document(candidate):
    return candidate

  raise HTTPException(status_code=404, detail="document not found")


def _html_document(title: str, body: str) -> str:
  safe_title = html.escape(title)
  return f"""
<!doctype html>
<html lang="it">
  <head>
    <meta charset="utf-8" />
    <title>{safe_title}</title>
    <style>
      html {{
        width: 100%;
        height: 100%;
        overflow: hidden;
      }}
      body {{
        margin: 0;
        padding: 10px;
        color: #1f2933;
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
        background: #fff;
        overflow: hidden;
      }}
      #fit-root {{
        transform-origin: top left;
      }}
      h1, h2 {{
        margin: 0 0 8px;
      }}
      h1 {{
        font-size: 15px;
      }}
      h2 {{
        margin-top: 12px;
        font-size: 13px;
      }}
      p {{
        margin: 0 0 8px;
        line-height: 1.35;
      }}
      table {{
        width: 100%;
        table-layout: fixed;
        border-collapse: collapse;
        margin: 0 0 10px;
        font-size: 12px;
        line-height: 1.2;
      }}
      td, th {{
        border: 1px solid #e7e0d6;
        padding: 3px 5px;
        vertical-align: top;
        overflow-wrap: anywhere;
        word-break: normal;
        white-space: pre-wrap;
      }}
    </style>
  </head>
  <body>
    <main id="fit-root">
      <h1>{safe_title}</h1>
      {body}
    </main>
    <script>
      function fitPreview() {{
        const root = document.getElementById('fit-root');
        if (!root) return;

        root.style.transform = 'none';
        root.style.width = 'auto';

        const availableWidth = Math.max(320, window.innerWidth - 4);
        const availableHeight = Math.max(240, window.innerHeight - 4);
        const contentWidth = Math.max(root.scrollWidth, root.offsetWidth);
        const contentHeight = Math.max(root.scrollHeight, root.offsetHeight);
        const scale = Math.min(
          1,
          availableWidth / contentWidth,
          availableHeight / contentHeight,
        );

        root.style.transform = `scale(${{scale}})`;
        root.style.width = `${{100 / scale}}%`;
      }}

      window.addEventListener('load', fitPreview);
      window.addEventListener('resize', fitPreview);
    </script>
  </body>
</html>
"""


def _excel_color(value) -> str | None:
  if not value:
    return None
  if getattr(value, "type", None) == "rgb" and value.rgb:
    rgb = str(value.rgb)
    if len(rgb) == 8:
      rgb = rgb[2:]
    return f"#{rgb}"
  return None


def _excel_border(side) -> str | None:
  if not side or not side.style:
    return None

  width = "2px" if side.style in {"medium", "thick", "double"} else "1px"
  color = _excel_color(side.color) or "#111111"
  return f"{width} solid {color}"


def _excel_cell_style(cell) -> str:
  rules: list[str] = []

  if cell.fill and cell.fill.fill_type:
    color = _excel_color(cell.fill.fgColor)
    if color and color != "#000000":
      rules.append(f"background:{color}")

  if cell.font:
    if cell.font.bold:
      rules.append("font-weight:700")
    if cell.font.italic:
      rules.append("font-style:italic")
    if cell.font.sz:
      rules.append(f"font-size:{float(cell.font.sz)}px")

  if cell.alignment:
    if cell.alignment.horizontal:
      horizontal = {
        "centerContinuous": "center",
        "general": "left",
      }.get(cell.alignment.horizontal, cell.alignment.horizontal)
      rules.append(f"text-align:{horizontal}")
    if cell.alignment.vertical:
      vertical = {
        "center": "middle",
        "top": "top",
        "bottom": "bottom",
      }.get(cell.alignment.vertical, cell.alignment.vertical)
      rules.append(f"vertical-align:{vertical}")

  border = cell.border
  if border:
    top = _excel_border(border.top)
    right = _excel_border(border.right)
    bottom = _excel_border(border.bottom)
    left = _excel_border(border.left)
    if top:
      rules.append(f"border-top:{top}")
    if right:
      rules.append(f"border-right:{right}")
    if bottom:
      rules.append(f"border-bottom:{bottom}")
    if left:
      rules.append(f"border-left:{left}")

  return ";".join(rules)


def _excel_cell_value(value) -> str:
  if value is None:
    return ""
  return html.escape(str(value))


def _excel_preview_html(path: Path) -> str:
  workbook = load_workbook(path, data_only=True)
  sections: list[str] = []

  for sheet in workbook.worksheets:
    merged_spans: dict[tuple[int, int], tuple[int, int]] = {}
    merged_children: set[tuple[int, int]] = set()
    for merged_range in sheet.merged_cells.ranges:
      min_col, min_row, max_col, max_row = merged_range.bounds
      merged_spans[(min_row, min_col)] = (
        max_row - min_row + 1,
        max_col - min_col + 1,
      )
      for row_index in range(min_row, max_row + 1):
        for col_index in range(min_col, max_col + 1):
          if row_index != min_row or col_index != min_col:
            merged_children.add((row_index, col_index))

    colgroup: list[str] = []
    for col_index in range(1, sheet.max_column + 1):
      letter = get_column_letter(col_index)
      width = sheet.column_dimensions[letter].width or 10
      pixel_width = max(36, min(180, int(width * 7)))
      colgroup.append(f'<col style="width:{pixel_width}px" />')

    rows: list[str] = []
    for row_index in range(1, sheet.max_row + 1):
      cells: list[str] = []
      for col_index in range(1, sheet.max_column + 1):
        if (row_index, col_index) in merged_children:
          continue

        cell = sheet.cell(row=row_index, column=col_index)
        rowspan, colspan = merged_spans.get((row_index, col_index), (1, 1))
        span_attrs = ""
        if rowspan > 1:
          span_attrs += f' rowspan="{rowspan}"'
        if colspan > 1:
          span_attrs += f' colspan="{colspan}"'

        style = _excel_cell_style(cell)
        style_attr = f' style="{style}"' if style else ""
        value = _excel_cell_value(cell.value)
        cells.append(f"<td{span_attrs}{style_attr}>{value}</td>")

      row_height = sheet.row_dimensions[row_index].height
      row_style = f' style="height:{row_height}px"' if row_height else ""
      rows.append(f"<tr{row_style}>{''.join(cells)}</tr>")

    if rows:
      sections.append(
        f"<h2>{html.escape(sheet.title)}</h2>"
        f"<table><colgroup>{''.join(colgroup)}</colgroup>{''.join(rows)}</table>"
      )

  if not sections:
    sections.append("<p>Il foglio Excel non contiene dati visualizzabili.</p>")

  return _html_document(path.name, "".join(sections))


def _docx_preview_html(path: Path) -> str:
  document = DocxDocument(path)
  parts: list[str] = []

  def paragraph_alignment(paragraph) -> str | None:
    alignment = paragraph.alignment
    if alignment is None and paragraph.style is not None:
      alignment = paragraph.style.paragraph_format.alignment
    if alignment is None:
      return None
    return {
      0: "left",
      1: "center",
      2: "right",
      3: "justify",
    }.get(int(alignment), None)

  def run_html(run) -> str:
    text = html.escape(run.text)
    if not text:
      return ""

    styles: list[str] = []
    if run.bold:
      styles.append("font-weight:700")
    if run.italic:
      styles.append("font-style:italic")
    if run.underline:
      styles.append("text-decoration:underline")
    if run.font.size:
      styles.append(f"font-size:{run.font.size.pt}pt")
    if run.font.color and run.font.color.rgb:
      styles.append(f"color:#{run.font.color.rgb}")

    style_attr = f' style="{";".join(styles)}"' if styles else ""
    return f"<span{style_attr}>{text}</span>"

  def paragraph_html(paragraph) -> str:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    text = paragraph.text.strip()
    content = "".join(run_html(run) for run in paragraph.runs) or html.escape(text)
    alignment = paragraph_alignment(paragraph)
    align_style = f' style="text-align:{alignment}"' if alignment else ""

    if style_name == "Title":
      return f'<h1 class="docx-title"{align_style}>{content}</h1>'
    if style_name.startswith("Heading"):
      return f'<h2 class="docx-heading"{align_style}>{content}</h2>'
    if text.startswith("•"):
      return f'<p class="docx-bullet"{align_style}>{content}</p>'
    return f'<p class="docx-paragraph"{align_style}>{content}</p>'

  for paragraph in document.paragraphs:
    if paragraph.text.strip():
      parts.append(paragraph_html(paragraph))
    else:
      parts.append('<div class="docx-spacer"></div>')

  for table in document.tables:
    rows: list[str] = []
    for row in table.rows:
      cells = "".join(
        f"<td>{html.escape(cell.text.strip())}</td>" for cell in row.cells
      )
      rows.append(f"<tr>{cells}</tr>")
    if rows:
      parts.append(f"<table>{''.join(rows)}</table>")

  if not parts:
    parts.append("<p>Il documento Word non contiene testo visualizzabile.</p>")

  safe_title = html.escape(path.name)
  return f"""
<!doctype html>
<html lang="it">
  <head>
    <meta charset="utf-8" />
    <title>{safe_title}</title>
    <style>
      html {{
        min-height: 100%;
        background: #f2f2f2;
      }}
      body {{
        margin: 0;
        padding: 28px;
        color: #000;
        background: #f2f2f2;
        font-family: Georgia, "Times New Roman", serif;
      }}
      .docx-page {{
        box-sizing: border-box;
        width: min(900px, calc(100vw - 56px));
        min-height: 1272px;
        margin: 0 auto;
        padding: 92px 104px;
        background: #fff;
        border-radius: 18px;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.14);
      }}
      .docx-title {{
        margin: 0 0 26px;
        padding-bottom: 8px;
        border-bottom: 2px solid #2f65ad;
        color: #173764;
        font-size: 31px;
        font-weight: 700;
        line-height: 1.1;
        text-transform: uppercase;
      }}
      .docx-heading {{
        margin: 22px 0 6px;
        color: #345f96;
        font-size: 24px;
        font-weight: 700;
        line-height: 1.2;
      }}
      .docx-paragraph,
      .docx-bullet {{
        margin: 0 0 9px;
        font-size: 20px;
        line-height: 1.23;
      }}
      .docx-bullet {{
        padding-left: 32px;
        text-indent: -18px;
      }}
      .docx-spacer {{
        height: 18px;
      }}
      table {{
        width: 100%;
        border-collapse: collapse;
        margin: 14px 0;
        font-size: 16px;
      }}
      td {{
        border: 1px solid #d1d5db;
        padding: 6px 8px;
        vertical-align: top;
      }}
      @media (max-width: 760px) {{
        body {{
          padding: 12px;
        }}
        .docx-page {{
          width: calc(100vw - 24px);
          padding: 48px 34px;
          border-radius: 12px;
        }}
        .docx-title {{
          font-size: 24px;
        }}
        .docx-heading {{
          font-size: 20px;
        }}
        .docx-paragraph,
        .docx-bullet {{
          font-size: 17px;
        }}
      }}
    </style>
  </head>
  <body>
    <main class="docx-page">
      {"".join(parts)}
    </main>
  </body>
</html>
"""


def _collect_documents(output_dir: Path, result: dict | None) -> list[Path]:
  documents: list[Path] = []
  if result:
    for path_str in result.get("documents_generated", []) or []:
      path = _resolve_path(str(path_str))
      if path.exists() and _is_visible_document(path):
        documents.append(path)

  if output_dir.exists():
    for path in sorted(output_dir.iterdir()):
      if path.is_file() and _is_visible_document(path) and path not in documents:
        documents.append(path)

  return documents


def _documents_for_job(job: JobRecord) -> list[Path]:
  documents = _collect_documents(job.output_dir, job.result if isinstance(job.result, dict) else None)
  if documents:
    job.documents = documents
  return documents


def _run_job(job_id: str) -> None:
  job = JOBS.get(job_id)
  if not job:
    return

  job.status = "running"
  job.step = "processing"

  try:
    os.chdir(BASE_DIR)
    get_config()

    job.output_dir.mkdir(parents=True, exist_ok=True)

    config = {"configurable": {"thread_id": f"job-{job_id}"}}
    initial_state = {
      "task": DEFAULT_TASK,
      "raw_document_path": str(job.file_path),
      "source_document_name": _original_upload_name(job.file_id, job.file_path),
      "job_id": job_id,
      "output_dir": str(job.output_dir),
      "messages": [HumanMessage(content=DEFAULT_TASK)],
    }
    result = workflow.invoke(initial_state, config)
    job.result = result
    job.documents = _collect_documents(job.output_dir, result if isinstance(result, dict) else None)
    if isinstance(result, dict) and result.get("ocr_review_level") == "error":
      job.status = "review_required"
      job.step = "review_required"
      job.completed_at = time.time()
      return

    job.status = "completed"
    job.step = "done"
    job.completed_at = time.time()
  except Exception as exc:
    job.status = "failed"
    job.step = "failed"
    job.error = str(exc)


def _run_generation_after_review(job_id: str) -> None:
  job = JOBS.get(job_id)
  if not job:
    return

  job.status = "running"
  job.step = "generating_documents"
  job.error = None

  try:
    os.chdir(BASE_DIR)
    get_config()

    previous_result = job.result if isinstance(job.result, dict) else {}
    resume_state = {
      **previous_result,
      "task": "Generate documents after human approval of fatal OCR findings.",
      "raw_document_path": str(job.file_path),
      "source_document_name": _original_upload_name(job.file_id, job.file_path),
      "job_id": job_id,
      "output_dir": str(job.output_dir),
      "resume_document_generation": True,
      "messages": [HumanMessage(content="Human approved document generation after OCR review.")],
    }
    config = {"configurable": {"thread_id": f"job-{job_id}-generation"}}
    result = workflow.invoke(resume_state, config)
    job.result = result
    job.documents = _collect_documents(job.output_dir, result if isinstance(result, dict) else None)
    job.status = "completed"
    job.step = "done"
    job.completed_at = time.time()
  except Exception as exc:
    job.status = "failed"
    job.step = "failed"
    job.error = str(exc)


@app.get("/health")
def health() -> dict[str, str]:
  return {"status": "ok"}


@app.post("/upload", response_model=UploadResponse)
async def upload(file: UploadFile = File(...)) -> UploadResponse:
  if not file.filename:
    raise HTTPException(status_code=400, detail="Missing filename")

  is_pdf = file.content_type == "application/pdf" or file.filename.lower().endswith(".pdf")
  if not is_pdf:
    raise HTTPException(status_code=400, detail="Only PDF files are supported")

  UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
  file_id = f"file_{uuid4().hex[:10]}"
  safe_name = _safe_filename(file.filename)
  destination = UPLOAD_DIR / f"{file_id}_{safe_name}"
  content = await file.read()
  if not content:
    raise HTTPException(status_code=400, detail="Empty file")

  destination.write_bytes(content)
  FILES[file_id] = destination
  return UploadResponse(fileId=file_id)


@app.post("/jobs/start", response_model=StartJobResponse)
def start_job(payload: StartJobRequest) -> StartJobResponse:
  file_path = FILES.get(payload.fileId)
  if not file_path:
    raise HTTPException(status_code=404, detail="fileId not found")

  job_id = f"job_{uuid4().hex[:10]}"
  output_dir = OUTPUT_DIR / job_id
  output_dir.mkdir(parents=True, exist_ok=True)
  JOBS[job_id] = JobRecord(
    file_id=payload.fileId,
    file_path=file_path,
    output_dir=output_dir,
    status="queued",
    step="queued",
  )

  worker = threading.Thread(target=_run_job, args=(job_id,), daemon=True)
  worker.start()

  return StartJobResponse(jobId=job_id)


@app.post("/jobs/{job_id}/continue-generation", response_model=ContinueGenerationResponse)
def continue_generation(job_id: str) -> ContinueGenerationResponse:
  job = JOBS.get(job_id)
  if not job:
    raise HTTPException(status_code=404, detail="jobId not found")
  if job.status != "review_required":
    raise HTTPException(
      status_code=409,
      detail="document generation can only be continued after a fatal OCR review block",
    )

  job.status = "running"
  job.step = "generating_documents"
  job.error = None
  worker = threading.Thread(target=_run_generation_after_review, args=(job_id,), daemon=True)
  worker.start()

  return ContinueGenerationResponse(jobId=job_id, status="running")


@app.get("/jobs/{job_id}", response_model=JobStatusResponse)
def get_job(job_id: str) -> JobStatusResponse:
  job = JOBS.get(job_id)
  if not job:
    raise HTTPException(status_code=404, detail="jobId not found")

  return JobStatusResponse(
    jobId=job_id,
    status=job.status,
    step=job.step,
    error=job.error,
  )


@app.get("/jobs/{job_id}/ocr-summary", response_model=OcrSummaryResponse)
def get_ocr_summary(job_id: str) -> OcrSummaryResponse:
  job = JOBS.get(job_id)
  if not job:
    raise HTTPException(status_code=404, detail="jobId not found")

  review_level: Literal["none", "warning", "error"] | None = None
  if job.result and isinstance(job.result, dict):
    review_level = job.result.get("ocr_review_level")

  text = ""
  ocr_path = job.output_dir / "ocr_output.txt"
  if ocr_path.exists():
    text = ocr_path.read_text(encoding="utf-8")

  if not text and job.result and isinstance(job.result, dict):
    text = str(job.result.get("extracted_data", "") or "")

  if not text:
    return OcrSummaryResponse(text="", reviewLevel=review_level)

  return OcrSummaryResponse(text=text, reviewLevel=review_level)


@app.get("/jobs/{job_id}/artifacts", response_model=list[Artifact])
def get_job_artifacts(job_id: str) -> list[Artifact]:
  job = JOBS.get(job_id)
  if not job:
    raise HTTPException(status_code=404, detail="jobId not found")

  if job.status != "completed":
    return []

  artifacts: list[Artifact] = []
  for path in _documents_for_job(job):
    doc_type, mime = _detect_type(path)
    download_url = f"/documents/{job_id}/{path.name}"
    preview_url = _preview_url(job_id, path, doc_type)
    artifacts.append(
      Artifact(
        id=f"{job_id}-{path.stem}",
        name=path.name,
        type=doc_type,
        mimeType=mime,
        previewUrl=preview_url,
        downloadUrl=download_url,
      )
    )

  return artifacts


@app.get("/documents", response_model=list[DocumentItem])
def list_documents() -> list[DocumentItem]:
  items: list[DocumentItem] = []
  for job_id, job in JOBS.items():
    if job.status != "completed":
      continue
    for path in _documents_for_job(job):
      doc_type, mime = _detect_type(path)
      download_url = f"/documents/{job_id}/{path.name}"
      preview_url = _preview_url(job_id, path, doc_type)
      created_at = _document_timestamp(path)
      items.append(
        DocumentItem(
          id=f"{job_id}-{path.stem}",
          name=path.name,
          type=doc_type,
          mimeType=mime,
          date=time.strftime("%Y-%m-%d %H:%M", time.localtime(created_at)),
          source="job",
          jobId=job_id,
          previewUrl=preview_url,
          downloadUrl=download_url,
        )
      )

  return items


@app.get("/documents/{job_id}/{filename}")
def download_document(job_id: str, filename: str):
  path = _find_document_path(job_id, filename)

  if path.suffix.lower() in {".txt", ".md"}:
    return PlainTextResponse(path.read_text(encoding="utf-8"))
  return FileResponse(path, filename=path.name)


@app.get("/documents/{job_id}/{filename}/preview", response_class=HTMLResponse)
def preview_document(job_id: str, filename: str) -> HTMLResponse:
  path = _find_document_path(job_id, filename)
  doc_type, _ = _detect_type(path)

  if doc_type == "excel":
    return HTMLResponse(_excel_preview_html(path))
  if doc_type == "word":
    return HTMLResponse(_docx_preview_html(path))

  raise HTTPException(status_code=415, detail="preview not supported")
